"""CG-07 data_source 家族：把父本里第一个字面量数值列表改成「从文件读」的各种写法（保型）。

候选的权威不变：仍然是 `父本 + 变换链 + seed`，数据文件由本模块**确定性**写出（同 seed 同父本逐字节相同），
读文件的表达式是固定模板。登记在 `compat_explore_transforms.TRANSFORMS`（family/axis 都是 `data_source`，
保型，同轴互斥，所以一条链里至多一个数据来源）。

读出的值必须与字面量一致：整数列表按整数读、含浮点的列表按浮点读（`dtype=` / `int()` / `float()` 显式写出），
否则刻度标签会变，保型检查会把模板的锅记成假差异。

两个由 seed 派生的轴（沿用 `_rng(seed, 父本, 变换名)`）：
* 路径写法：脚本目录相对（`os.path.dirname(__file__)`）/ cwd 相对 / `Path(__file__).resolve().parent / "data"`；
* 编码（仅文本格式）：utf-8 / utf-8-sig / gbk / 中文文件名 `实验数据`。读端一律显式写 encoding。

二进制格式的确定性：npy 头固定；pickle 固定 protocol 4；xlsx 由 openpyxl 写出后重新打包（固定时间戳、
ZIP_STORED，不依赖 zlib 实现，并抹掉 docProps 里的 created/modified）；parquet 不压缩、不写统计、
不带 pandas 元数据（footer 里的 `created_by` 带 pyarrow 版本，所以版本在 workflow 里钉死）。

可选依赖（openpyxl / pyarrow 只装进 Guardian workflow，产品依赖不变）缺失时，对应变换 `applies` 为 False，
原因经 `unavailable()` 进 coverage。这个模块只在 explore 进程里运行，不碰产品代码。
"""

from __future__ import annotations

import ast
import importlib.util
import io
import re
import zipfile
from dataclasses import dataclass
from pathlib import Path

#: (读端 encoding, 文件名主干)。一项一个取值，不是两个独立轴：中文文件名与 utf-8 搭配。
ENCODINGS = (
    ("utf-8", "cg_data"),
    ("utf-8-sig", "cg_data"),
    ("gbk", "cg_data"),
    ("utf-8", "实验数据"),
)
PATH_STYLES = ("script_dir", "cwd", "data_subdir")
#: 文本格式的表头。第一列是中文，gbk / utf-8-sig 才真的改变字节。
HEADER = ("序号", "y")
ZIP_EPOCH = (1980, 1, 1, 0, 0, 0)

_TSV_HELPER = """
def _cg_read_tsv(path, encoding):
    with open(path, newline="", encoding=encoding) as fh:
        return [CONV(row["y"]) for row in _cg_csv.DictReader(fh, delimiter="\\t")]
"""
_PARTS_HELPER = """
def _cg_read_parts(pattern, encoding):
    values = []
    for part in sorted(_cg_glob.glob(pattern)):
        with open(part, newline="", encoding=encoding) as fh:
            values.extend(CONV(row["y"]) for row in _cg_csv.DictReader(fh))
    return values
"""


def _xt():
    import compat_explore_transforms as XT

    return XT


def _target(ctx):
    XT = _xt()
    found = XT._inline_target(ctx)
    if found is None:
        return None
    call, i = found
    return call, i, XT._numeric_list(call.args[i])


def _conv(values: list) -> str:
    """读回的元素类型要和字面量一致：含浮点按浮点，否则按整数。"""
    return "float" if any(type(v) is float for v in values) else "int"


# ---------------------------------------------------------------- writers (bytes, deterministic)


def _text(values: list, encoding: str, delimiter: str = ",", start: int = 0) -> bytes:
    lines = [delimiter.join(HEADER)]
    lines += [f"{start + i}{delimiter}{v!r}" for i, v in enumerate(values)]
    return ("\n".join(lines) + "\n").encode(encoding)


def _numeric_dtype(values: list):
    import numpy as np

    return np.float64 if _conv(values) == "float" else np.int64


def _write_npy(values: list, encoding) -> bytes:
    import numpy as np

    buf = io.BytesIO()
    np.save(buf, np.array(values, dtype=_numeric_dtype(values)), allow_pickle=False)
    return buf.getvalue()


def _write_pickle(values: list, encoding) -> bytes:
    import pickle

    import numpy as np
    import pandas as pd

    frame = pd.DataFrame({"y": np.array(values, dtype=_numeric_dtype(values))})
    return pickle.dumps(frame, protocol=4)


def _repack_fixed(raw: bytes) -> bytes:
    """xlsx 是 zip：重新打包成逐字节确定（固定时间、顺序、不压缩），并抹掉 core.xml 里的写入时刻。"""
    out = io.BytesIO()
    with (
        zipfile.ZipFile(io.BytesIO(raw)) as src,
        zipfile.ZipFile(out, "w", zipfile.ZIP_STORED) as dst,
    ):
        for name in sorted(src.namelist()):
            data = src.read(name)
            if name == "docProps/core.xml":
                data = re.sub(
                    rb"(<dcterms:(created|modified)[^>]*>)[^<]*(</dcterms:\2>)",
                    rb"\g<1>2000-01-01T00:00:00Z\g<3>",
                    data,
                )
            info = zipfile.ZipInfo(name, ZIP_EPOCH)
            info.compress_type = zipfile.ZIP_STORED
            info.create_system = 3
            info.external_attr = 0o600 << 16
            dst.writestr(info, data)
    return out.getvalue()


def _write_xlsx(values: list, encoding) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet"
    ws.append(["y"])
    for v in values:
        ws.append([v])
    wb.properties.creator = "cg"
    buf = io.BytesIO()
    wb.save(buf)
    return _repack_fixed(buf.getvalue())


def _write_parquet(values: list, encoding) -> bytes:
    import pyarrow as pa
    import pyarrow.parquet as pq

    buf = io.BytesIO()
    table = pa.table(
        {"y": pa.array(values, type=pa.float64() if _conv(values) == "float" else pa.int64())}
    )
    pq.write_table(
        table,
        buf,
        compression="NONE",
        use_dictionary=False,
        write_statistics=False,
        store_schema=False,
    )
    return buf.getvalue()


# ---------------------------------------------------------------- formats


@dataclass(frozen=True)
class Format:
    name: str
    order: int
    requires: tuple[str, ...]
    ext: str
    text: bool
    imports: tuple[str, ...]
    summary: str
    write: object = None  # None = text (see _text)
    expr: object = None  # (path_src, encoding, conv) -> expression source
    helper: str | None = None


def _e(enc) -> str:
    return repr(enc)


FORMATS = (
    Format(
        "read_csv_pandas",
        21,
        ("pandas",),
        ".csv",
        True,
        ("import pandas as _cg_pd",),
        "read the first literal data list with pandas.read_csv(...)['y'].tolist()",
        expr=lambda p, enc, k: (
            f'_cg_pd.read_csv({p}, encoding={_e(enc)}, float_precision="round_trip")["y"].tolist()'
        ),
    ),
    Format(
        "read_csv_numpy",
        22,
        ("numpy",),
        ".csv",
        True,
        ("import numpy as _cg_np",),
        "read the first literal data list with numpy.loadtxt(..., skiprows=1)",
        expr=lambda p, enc, k: (
            f'_cg_np.loadtxt({p}, delimiter=",", skiprows=1, usecols=1, dtype={k}, encoding={_e(enc)})'
        ),
    ),
    Format(
        "genfromtxt_names",
        23,
        ("numpy",),
        ".csv",
        True,
        ("import numpy as _cg_np",),
        "read the first literal data list with numpy.genfromtxt(..., names=True)['y']",
        expr=lambda p, enc, k: (
            f'_cg_np.genfromtxt({p}, names=True, delimiter=",", dtype={k}, encoding={_e(enc)})["y"]'
        ),
    ),
    Format(
        "read_tsv_stdlib",
        24,
        (),
        ".tsv",
        True,
        ("import csv as _cg_csv",),
        "read the first literal data list from a TSV with the csv module (newline='', explicit encoding)",
        expr=lambda p, enc, k: f"_cg_read_tsv({p}, {_e(enc)})",
        helper=_TSV_HELPER,
    ),
    Format(
        "read_npy",
        25,
        ("numpy",),
        ".npy",
        False,
        ("import numpy as _cg_np",),
        "read the first literal data list with numpy.load of a .npy file",
        write=_write_npy,
        expr=lambda p, enc, k: f"_cg_np.load({p})",
    ),
    Format(
        "read_pickle",
        26,
        ("pandas", "numpy"),
        ".pkl",
        False,
        ("import pandas as _cg_pd",),
        "read the first literal data list with pandas.read_pickle(...)['y'].tolist()",
        write=_write_pickle,
        expr=lambda p, enc, k: f'_cg_pd.read_pickle({p})["y"].tolist()',
    ),
    Format(
        "read_excel",
        27,
        ("openpyxl",),
        ".xlsx",
        False,
        ("import openpyxl as _cg_openpyxl",),
        "read the first literal data list from an .xlsx with openpyxl.load_workbook",
        write=_write_xlsx,
        expr=lambda p, enc, k: (
            f"[{k}(row[0]) for row in "
            f"_cg_openpyxl.load_workbook({p}).active.iter_rows(min_row=2, values_only=True)]"
        ),
    ),
    Format(
        "read_parquet",
        28,
        ("pyarrow",),
        ".parquet",
        False,
        ("import pyarrow.parquet as _cg_pq",),
        "read the first literal data list with pyarrow.parquet.read_table",
        write=_write_parquet,
        expr=lambda p, enc, k: f'_cg_pq.read_table({p}).column("y").to_pylist()',
    ),
    Format(
        "glob_dir",
        29,
        (),
        ".csv",
        True,
        ("import csv as _cg_csv", "import glob as _cg_glob"),
        "split the first literal data list over data/*.csv and read it back with sorted(glob(...))",
        expr=lambda p, enc, k: f"_cg_read_parts({p}, {_e(enc)})",
        helper=_PARTS_HELPER,
    ),
)
BY_NAME = {f.name: f for f in FORMATS}


def missing(fmt: Format) -> list[str]:
    return [m for m in fmt.requires if importlib.util.find_spec(m) is None]


def unavailable() -> dict[str, str]:
    """格式 → 缺的模块。给 coverage 用：缺依赖的变换不适用，原因要说出来。"""
    return {f.name: "missing module(s): " + ", ".join(missing(f)) for f in FORMATS if missing(f)}


# ---------------------------------------------------------------- apply


def _path_src(style: str, subdir: str, name: str) -> tuple[str, tuple[str, ...]]:
    """(路径表达式源码, 需要的 import)。subdir 为空时文件与脚本同目录。"""
    parts = [repr(p) for p in (subdir, name) if p]
    if style == "script_dir":
        return (
            f"_cg_os.path.join(_cg_os.path.dirname(__file__), {', '.join(parts)})",
            ("import os as _cg_os",),
        )
    if style == "cwd":
        return repr("/".join(p for p in (subdir, name) if p)), ()
    return (
        f"str(_cg_Path(__file__).resolve().parent / {' / '.join(parts)})",
        ("from pathlib import Path as _cg_Path",),
    )


def _add_helper(ctx, source: str, conv: str) -> None:
    """在文件头的 import 段后面放一个模块级 helper（import 都在它前面）。"""
    node = ast.parse(source.replace("CONV", conv)).body[0]
    body = ctx.tree.body
    at = 0
    for i, stmt in enumerate(body):
        if isinstance(stmt, (ast.Import, ast.ImportFrom)) or (i == 0 and _xt()._is_docstring(stmt)):
            at = i + 1
        else:
            break
    body.insert(at, node)


def _chunks(values: list, count: int) -> list[list]:
    base, extra = divmod(len(values), count)
    out, at = [], 0
    for i in range(count):
        size = base + (1 if i < extra else 0)
        out.append(values[at : at + size])
        at += size
    return out


def _applies(fmt: Format):
    def applies(ctx) -> bool:
        return not missing(fmt) and _target(ctx) is not None

    return applies


def _apply(fmt: Format):
    def apply(ctx, rng) -> None:
        XT = _xt()
        call, i, values = _target(ctx)
        style = rng.choice(PATH_STYLES)
        enc, base = rng.choice(ENCODINGS) if fmt.text else (None, "cg_data")
        conv = _conv(values)
        folder = Path(ctx.script).parent
        glob_dir = fmt.name == "glob_dir"
        subdir = "data" if style == "data_subdir" or glob_dir else ""
        name = f"{base}_*{fmt.ext}" if glob_dir else f"{base}{fmt.ext}"
        written: list[str] = []
        if glob_dir:
            parts = _chunks(values, min(len(values), 2 + rng.randrange(2)))
            start = 0
            for n, part in enumerate(parts):
                rel = (folder / subdir / f"{base}_{n:02d}{fmt.ext}").as_posix()
                ctx.files[rel] = _text(part, enc, start=start)
                written.append(rel)
                start += len(part)
        else:
            rel = (folder / subdir / name).as_posix()
            if fmt.write is not None:
                ctx.files[rel] = fmt.write(values, enc)
            else:
                delimiter = "\t" if fmt.ext == ".tsv" else ","
                ctx.files[rel] = _text(values, enc, delimiter)
            written.append(rel)
        path_src, path_imports = _path_src(style, subdir, name)
        expr = fmt.expr(path_src, enc, conv)
        call.args[i] = ast.parse(expr, mode="eval").body
        if fmt.helper:
            _add_helper(ctx, fmt.helper, conv)
        for line in (*fmt.imports, *path_imports):
            XT._insert_import(ctx.tree, line)
        ctx.params.update(
            {
                "format": fmt.name,
                "path_style": style,
                "encoding": enc,
                "data_files": written,
                "read_expr": expr,
            }
        )

    return apply


@dataclass(frozen=True)
class Spec:
    name: str
    order: int
    applies: object
    apply: object
    summary: str


SPECS = tuple(Spec(f.name, f.order, _applies(f), _apply(f), f.summary) for f in FORMATS)
